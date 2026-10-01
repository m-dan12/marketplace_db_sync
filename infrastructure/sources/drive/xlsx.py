"""Turns the nightly-export xlsx files (Drive archive) back into the same
domain lines the API sources produce, so a backfilled row is
indistinguishable from a live one. Pure functions over `list[dict]` rows —
no Drive/Postgres here, which keeps them unit-testable."""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any, Iterable, Optional

from openpyxl import load_workbook

from domain.models import (
    OzonOrderLine,
    OzonPriceLine,
    OzonStockLine,
    OzonWarehouseStockLine,
    SelsupMovementLine,
    SelsupStockLine,
    WbAdStatLine,
    WbOrderLine,
    WbPriceLine,
    WbSaleLine,
    WbStockLine,
)
from infrastructure.sources.parsing import parse_date, parse_datetime

Row = dict[str, Any]

# Columns of the WB stock report that are not physical warehouses' own data
# are still stored as "warehouses" by the live API source (WB returns them
# in the same `warehouses` array), so the backfill keeps them for parity.
_WB_STOCK_FIXED_COLUMNS = ("nmId", "vendorCode", "barcode", "techSize", "volume")


def read_rows(data: bytes, sheet: Optional[str] = None) -> list[Row]:
    """First sheet (or the named one) as a list of dicts keyed by header."""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        iterator = worksheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            return []
        names = [str(h).strip() if h is not None else "" for h in header]
        rows: list[Row] = []
        for values in iterator:
            if values is None or all(v is None for v in values):
                continue
            rows.append({names[i]: values[i] for i in range(min(len(names), len(values))) if names[i]})
        return rows
    finally:
        workbook.close()


def sheet_names(data: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


# --- scalar helpers ---------------------------------------------------------

def to_str(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].lstrip("-").isdigit():
        text = text[:-2]
    return text or None


def to_int(value: Any) -> Optional[int]:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def to_float(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _as_date(value: Any) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return parse_date(to_str(value))


def _as_datetime(value: Any) -> Optional[datetime]:
    if isinstance(value, datetime):
        return value
    return parse_datetime(to_str(value))


def _naive(value: Optional[datetime]) -> Optional[datetime]:
    return value.replace(tzinfo=None) if value is not None else None


# --- WB ---------------------------------------------------------------------

def parse_wb_orders(rows: Iterable[Row]) -> list[WbOrderLine]:
    return [
        WbOrderLine(
            srid=to_str(r["srid"]),
            order_date=_as_date(r.get("date")),
            last_change_date=_as_datetime(r.get("lastChangeDate")),
            warehouse_name=to_str(r.get("warehouseName")),
            region_name=to_str(r.get("regionName")),
            supplier_article=to_str(r.get("supplierArticle")),
            nm_id=to_int(r.get("nmId")),
            barcode=to_str(r.get("barcode")),
            subject=to_str(r.get("subject")),
            brand=to_str(r.get("brand")),
            tech_size=to_str(r.get("techSize")),
            total_price=to_float(r.get("totalPrice")),
            discount_percent=to_float(r.get("discountPercent")),
            finished_price=to_float(r.get("finishedPrice")),
            price_with_disc=to_float(r.get("priceWithDisc")),
            is_cancel=_as_bool(r.get("isCancel")),
            g_number=to_str(r.get("gNumber")),
        )
        for r in rows
        if r.get("srid")
    ]


def parse_wb_sales(rows: Iterable[Row]) -> list[WbSaleLine]:
    return [
        WbSaleLine(
            sale_id=to_str(r["saleID"]),
            sale_date=_as_date(r.get("date")),
            last_change_date=_as_datetime(r.get("lastChangeDate")),
            warehouse_name=to_str(r.get("warehouseName")),
            region_name=to_str(r.get("regionName")),
            supplier_article=to_str(r.get("supplierArticle")),
            nm_id=to_int(r.get("nmId")),
            barcode=to_str(r.get("barcode")),
            subject=to_str(r.get("subject")),
            brand=to_str(r.get("brand")),
            tech_size=to_str(r.get("techSize")),
            total_price=to_float(r.get("totalPrice")),
            discount_percent=to_float(r.get("discountPercent")),
            spp=to_float(r.get("spp")),
            for_pay=to_float(r.get("forPay")),
            finished_price=to_float(r.get("finishedPrice")),
            price_with_disc=to_float(r.get("priceWithDisc")),
            order_type=to_str(r.get("orderType")),
            g_number=to_str(r.get("gNumber")),
        )
        for r in rows
        if r.get("saleID")
    ]


def _as_bool(value: Any) -> bool:
    if isinstance(value, str):
        return value.strip().lower() in ("true", "1", "да", "yes")
    return bool(value)


def parse_wb_stocks(rows: Iterable[Row]) -> list[WbStockLine]:
    """Wide report (one column per warehouse) -> one line per non-zero
    (barcode, warehouse)."""
    lines: list[WbStockLine] = []
    for r in rows:
        nm_id = to_int(r.get("nmId"))
        if nm_id is None:
            continue
        for column, value in r.items():
            if column in _WB_STOCK_FIXED_COLUMNS:
                continue
            quantity = to_float(value)
            if not quantity:
                continue
            lines.append(
                WbStockLine(
                    nm_id=nm_id,
                    vendor_code=to_str(r.get("vendorCode")),
                    barcode=to_str(r.get("barcode")),
                    tech_size=to_str(r.get("techSize")),
                    volume=to_float(r.get("volume")),
                    warehouse_name=column,
                    quantity=quantity,
                )
            )
    return lines


def parse_wb_prices(rows: Iterable[Row]) -> list[WbPriceLine]:
    return [
        WbPriceLine(
            nm_id=to_int(r["nmID"]),
            vendor_code=to_str(r.get("vendorCode")),
            tech_size=to_str(r.get("techSizeName")),
            size_id=to_int(r.get("sizeID")),
            price=to_float(r.get("price")),
            discount=to_float(r.get("discount")),
            discounted_price=to_float(r.get("discountedPrice")),
        )
        for r in rows
        if to_int(r.get("nmID")) is not None
    ]


def parse_wb_ads(rows: Iterable[Row]) -> list[WbAdStatLine]:
    """Sheet `По товарам` of the `Реклама` export."""
    return [
        WbAdStatLine(
            stat_date=_as_date(r.get("Дата")),
            campaign_id=to_int(r["ID кампании"]),
            nm_id=to_int(r["nmId"]),
            campaign_status=to_str(r.get("Статус")),
            views=to_int(r.get("Показы")),
            clicks=to_int(r.get("Клики")),
            ctr=to_float(r.get("CTR%")),
            cpc=to_float(r.get("CPC")),
            spend=to_float(r.get("Затраты")),
            orders=to_int(r.get("Заказы")),
            carts=to_int(r.get("Корзина")),
            avg_position=to_float(r.get("Ср.позиция")),
        )
        for r in rows
        if r.get("ID кампании") is not None and r.get("nmId") is not None and _as_date(r.get("Дата"))
    ]


# --- Ozon -------------------------------------------------------------------

def parse_ozon_orders(rows: Iterable[Row]) -> list[OzonOrderLine]:
    return [
        OzonOrderLine(
            posting_number=to_str(r["posting_number"]),
            offer_id=to_str(r["offer_id"]),
            status=to_str(r.get("status")),
            order_date=_as_datetime(r.get("order_date")),
            source=to_str(r.get("source")),
            warehouse_name=to_str(r.get("склад_отгрузки")),
            city=to_str(r.get("город_доставки")),
            region=to_str(r.get("регион")),
            sku=to_int(r.get("sku")),
            product_name=to_str(r.get("product_name")),
            quantity=to_int(r.get("quantity")),
            price=to_float(r.get("price")),
            currency=to_str(r.get("currency")),
        )
        for r in rows
        if r.get("posting_number") and r.get("offer_id")
    ]


def parse_ozon_stocks(rows: Iterable[Row]) -> list[OzonStockLine]:
    return [
        OzonStockLine(
            offer_id=to_str(r["offer_id"]),
            product_id=to_int(r.get("product_id")),
            sku=to_int(r.get("sku")),
            stock_type=to_str(r.get("type")),
            present=to_int(r.get("present")),
            reserved=to_int(r.get("reserved")),
        )
        for r in rows
        if r.get("offer_id") and r.get("type")
    ]


def parse_ozon_prices(rows: Iterable[Row]) -> list[OzonPriceLine]:
    return [
        OzonPriceLine(
            offer_id=to_str(r["offer_id"]),
            product_id=to_int(r.get("product_id")),
            price=to_float(r.get("price")),
            old_price=to_float(r.get("old_price")),
            min_price=to_float(r.get("min_price")),
            marketing_seller_price=to_float(r.get("marketing_seller_price")),
        )
        for r in rows
        if r.get("offer_id")
    ]


def parse_ozon_warehouse_stocks(rows: Iterable[Row]) -> list[OzonWarehouseStockLine]:
    """Sheet `Остатки по складам FBO` of the Ozon stock export."""
    return [
        OzonWarehouseStockLine(
            offer_id=to_str(r["offer_id"]),
            sku=to_int(r.get("sku")),
            product_name=to_str(r.get("товар")),
            warehouse_name=to_str(r["склад"]),
            free_to_sell=to_int(r.get("доступно")),
            reserved=to_int(r.get("резерв")),
            promised=to_int(r.get("в пути")),
        )
        for r in rows
        if r.get("offer_id") and r.get("склад")
    ]


# --- Selsup -----------------------------------------------------------------

def sku_accounts_from_rows(
    rows: Iterable[Row], organization_accounts: dict[int, str]
) -> dict[int, str]:
    """`skuId -> account` from a stock file that has `organizationId`."""
    mapping: dict[int, str] = {}
    for r in rows:
        account = organization_accounts.get(to_int(r.get("organizationId")))
        sku_id = to_int(r.get("skuId"))
        if account is not None and sku_id is not None:
            mapping[sku_id] = account
    return mapping


def parse_selsup_stocks(
    rows: Iterable[Row],
    organization_accounts: dict[int, str],
    sku_accounts: Optional[dict[int, str]] = None,
) -> dict[str, list[SelsupStockLine]]:
    """Split by account through `organizationId`; zero-quantity rows are
    skipped, as in the live source. The early daily files (before 21.08.2026)
    have no `organizationId` column: their rows are attributed through
    `sku_accounts` (skuId -> account taken from a newer file), and rows of
    SKUs unknown to it are dropped."""
    by_account: dict[str, list[SelsupStockLine]] = {}
    for r in rows:
        if not to_float(r.get("quantity")):
            continue
        if "organizationId" in r:
            account = organization_accounts.get(to_int(r.get("organizationId")))
        else:
            account = (sku_accounts or {}).get(to_int(r.get("skuId")))
        if account is None:
            continue
        by_account.setdefault(account, []).append(
            SelsupStockLine(
                warehouse_id=to_int(r["warehouseId"]),
                warehouse_name=to_str(r.get("warehouseName")),
                sku_id=to_int(r["skuId"]),
                article=to_str(r.get("article")),
                wb_size=to_str(r.get("wbSize")),
                ozon_article=to_str(r.get("ozonArticle")),
                cell_name=to_str(r.get("cellName")),
                quantity=to_float(r.get("quantity")),
                available_quantity=to_float(r.get("availableQuantity")),
                calculated_quantity=to_float(r.get("calculatedQuantity")),
                modify_date=_as_datetime(r.get("modifyDate")),
            )
        )
    return by_account


def parse_selsup_movements(rows: Iterable[Row]) -> list[SelsupMovementLine]:
    lines: list[SelsupMovementLine] = []
    for r in rows:
        moved_at = _naive(_as_datetime(r.get("date")))
        sku_id = to_int(r.get("skuId"))
        quantity = to_float(r.get("quantity"))
        if moved_at is None or sku_id is None or quantity is None:
            continue
        lines.append(
            SelsupMovementLine(
                movement_type=to_str(r.get("тип")),
                operation=to_str(r.get("operation")),
                moved_at=moved_at,
                warehouse_id=to_int(r.get("warehouseId")),
                order_id=to_int(r.get("orderId")),
                order_type=to_str(r.get("orderType")),
                sku_id=sku_id,
                article=to_str(r.get("артикул")),
                product_name=to_str(r.get("название")),
                cell_name=to_str(r.get("cellName")),
                quantity=quantity,
                user_id=to_int(r.get("userId")),
            )
        )
    return lines
